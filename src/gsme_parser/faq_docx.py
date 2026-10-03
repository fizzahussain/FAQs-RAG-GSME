from pathlib import Path
import re

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from gsme_parser.models import FAQRecord, ParseResult, ParseWarning
from gsme_parser.text import (
    make_faq_id,
    normalize_inline,
    normalize_multiline,
    split_question_number,
)


def _paragraph_text(paragraph: Paragraph) -> str:
    parts = []
    for node in paragraph._p.iter(qn("w:t")):
        if node.text:
            parts.append(node.text)
    return normalize_inline("".join(parts))


def _cell_text(cell) -> str:
    parts = [_paragraph_text(paragraph) for paragraph in cell.paragraphs]
    return normalize_multiline("\n\n".join(part for part in parts if part))


def _bold_share(paragraph: Paragraph) -> float:
    total = 0
    bold = 0
    for run in paragraph.runs:
        text = run.text.strip()
        if not text:
            continue
        total += len(text)
        if run.bold is True:
            bold += len(text)
    return bold / total if total else 0.0


def _is_qa_table(table: Table) -> bool:
    if not table.rows or len(table.columns) < 2:
        return False
    first = [normalize_inline(cell.text).lower() for cell in table.rows[0].cells[:2]]
    return first == ["question", "answer"]


def _parse_table(path: Path, table: Table, table_number: int) -> ParseResult:
    result = ParseResult()
    category = "Uncategorized"
    seen_ids: set[str] = set()

    for row_number, row in enumerate(table.rows[1:], start=2):
        question_cell = _cell_text(row.cells[0])
        answer_cell = _cell_text(row.cells[1])
        location = f"table {table_number}, row {row_number}"

        if not question_cell and not answer_cell:
            continue

        if (
            question_cell
            and answer_cell
            and normalize_inline(question_cell).casefold()
            == normalize_inline(answer_cell).casefold()
            and "?" not in question_cell
        ):
            category = normalize_inline(question_cell)
            continue

        number, question = split_question_number(question_cell)
        answer = normalize_multiline(answer_cell)

        if not question or not answer:
            result.warnings.append(
                ParseWarning(
                    code="incomplete_faq_row",
                    message="FAQ row is missing a question or answer",
                    location=location,
                )
            )
            continue

        faq_id = make_faq_id(category, number, question)
        if faq_id in seen_ids:
            result.warnings.append(
                ParseWarning(
                    code="duplicate_faq_id",
                    message=f"Duplicate FAQ identifier {faq_id}",
                    location=location,
                )
            )
            faq_id = make_faq_id(category, None, question)
        seen_ids.add(faq_id)

        result.faq_records.append(
            FAQRecord(
                faq_id=faq_id,
                faq_number=number,
                category=category,
                question=question,
                answer=answer,
                source_file=path.name,
                source_kind="docx_table",
                source_location=location,
            )
        )

    return result


def _looks_like_category(text: str, style_name: str, next_text: str | None) -> bool:
    if style_name.lower().startswith("heading"):
        return True
    if len(text) > 80 or "?" in text or text.endswith((".", ":", ";", "!")):
        return False
    if re.fullmatch(r"\d{1,3}", text):
        return False
    return bool(next_text and re.fullmatch(r"\d{1,3}", next_text))


def _parse_paragraphs(path: Path, document: DocumentObject) -> ParseResult:
    result = ParseResult()
    paragraphs = [p for p in document.paragraphs if _paragraph_text(p)]
    category = "Uncategorized"
    pending_number: str | None = None
    current_question: str | None = None
    current_number: str | None = None
    current_answer: list[str] = []
    current_location: str | None = None
    seen_ids: set[str] = set()

    def flush() -> None:
        nonlocal current_question, current_number, current_answer, current_location
        if current_question is None:
            return
        answer = normalize_multiline("\n\n".join(current_answer))
        if not answer:
            result.warnings.append(
                ParseWarning(
                    code="missing_answer",
                    message=f"No answer found for question: {current_question}",
                    location=current_location,
                )
            )
        else:
            faq_id = make_faq_id(category, current_number, current_question)
            if faq_id in seen_ids:
                faq_id = make_faq_id(category, None, current_question)
            seen_ids.add(faq_id)
            result.faq_records.append(
                FAQRecord(
                    faq_id=faq_id,
                    faq_number=current_number,
                    category=category,
                    question=current_question,
                    answer=answer,
                    source_file=path.name,
                    source_kind="docx_paragraphs",
                    source_location=current_location or "paragraph",
                )
            )
        current_question = None
        current_number = None
        current_answer = []
        current_location = None

    for index, paragraph in enumerate(paragraphs):
        text = _paragraph_text(paragraph)
        style_name = paragraph.style.name if paragraph.style else "Normal"
        next_text = _paragraph_text(paragraphs[index + 1]) if index + 1 < len(paragraphs) else None
        location = f"paragraph {index + 1}"

        if _looks_like_category(text, style_name, next_text):
            flush()
            category = text
            pending_number = None
            continue

        if re.fullmatch(r"\d{1,3}", text):
            flush()
            pending_number = text.zfill(2)
            continue

        number, stripped_question = split_question_number(text)
        question_signal = (
            text.endswith("?")
            and (
                number is not None
                or pending_number is not None
                or _bold_share(paragraph) >= 0.6
            )
        )

        if question_signal:
            flush()
            current_question = stripped_question
            current_number = number or pending_number
            current_location = location
            pending_number = None
            continue

        if current_question is not None:
            current_answer.append(text)

    flush()

    if not result.faq_records:
        result.warnings.append(
            ParseWarning(
                code="no_faq_records",
                message="No FAQ records could be identified from DOCX paragraphs",
            )
        )
    return result


def parse_faq_docx(path: str | Path, strategy: str = "auto") -> ParseResult:
    input_path = Path(path)
    document = Document(input_path)

    if strategy not in {"auto", "table", "paragraphs"}:
        raise ValueError("DOCX strategy must be auto, table, or paragraphs")

    if strategy in {"auto", "table"}:
        for table_number, table in enumerate(document.tables, start=1):
            if _is_qa_table(table):
                return _parse_table(input_path, table, table_number)
        if strategy == "table":
            return ParseResult(
                warnings=[
                    ParseWarning(
                        code="qa_table_not_found",
                        message="No table with Question and Answer headers was found",
                    )
                ]
            )

    return _parse_paragraphs(input_path, document)
