from collections.abc import Iterator
import hashlib
from pathlib import Path
import re

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.hyperlink import Hyperlink

from gsme_rag.parsing.markdown import finish_markdown, render_markdown_table
from gsme_rag.parsing.ocr import OcrUnavailable, ocr_image
from gsme_rag.models import ParseWarning, ParsedDocument
from gsme_rag.text import normalize_inline, normalize_multiline
from gsme_rag.parsing.fidelity import _reason as fidelity_reason


def _iter_blocks(document: DocumentObject) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _format_runs(paragraph: Paragraph, include_emphasis: bool = True) -> str:
    groups: list[tuple[str, bool, bool, str]] = []
    # Hyperlinks contain their own runs and are absent from paragraph.runs.
    runs = [run for item in paragraph.iter_inner_content()
            for run in (item.runs if isinstance(item, Hyperlink) else [item])]
    for run in runs:
        text = run.text.replace("\r", "").replace("\n", " ")
        if not text:
            continue
        bold = next((value for value in (run.bold, run.style.font.bold, paragraph.style.font.bold)
                     if value is not None), False)
        italic = next((value for value in (run.italic, run.style.font.italic, paragraph.style.font.italic)
                       if value is not None), False)
        position = "sup" if run.font.superscript else "sub" if run.font.subscript else ""
        if groups and groups[-1][1:] == (bold, italic, position):
            previous = groups[-1][0]
            groups[-1] = (previous + text, bold, italic, position)
        else:
            groups.append((text, bold, italic, position))

    fragments = []
    for text, bold, italic, position in groups:
        if position:
            text = f"<{position}>{text}</{position}>"
        marker = ("***" if bold and italic else "**" if bold else "*" if italic else "") if include_emphasis else ""
        fragments.append(f"{marker}{text}{marker}" if marker else text)
    rendered = "".join(fragments)
    return re.sub(r"[ \t]+", " ", rendered).strip()


def _paragraph_to_markdown(paragraph: Paragraph) -> str:
    style = paragraph.style.name if paragraph.style else "Normal"
    if style.casefold() in {"code", "source code", "html preformatted"}:
        return "```\n" + paragraph.text.rstrip("\r\n") + "\n```"
    text = _format_runs(paragraph)
    if not text:
        return ""

    heading_match = re.match(r"Heading\s+([1-6])", style, flags=re.IGNORECASE)
    if heading_match:
        level = int(heading_match.group(1))
        return f"{'#' * level} {normalize_inline(_format_runs(paragraph, include_emphasis=False))}"

    if style.lower().startswith("title"):
        return f"# {normalize_inline(_format_runs(paragraph, include_emphasis=False))}"
    if "list bullet" in style.lower():
        return f"- {text}"
    if "list number" in style.lower():
        return f"1. {text}"
    return text


def _image_ids(element) -> Iterator[str]:
    for node in element.iter():
        if node.tag == qn("a:blip"):
            value = node.get(qn("r:embed"))
        elif node.tag == "{urn:schemas-microsoft-com:vml}imagedata":
            value = node.get(qn("r:id"))
        else:
            continue
        if value:
            yield value


def _image_text(element, part, location: str, warnings: list[ParseWarning],
                cache: dict[str, str | ValueError]) -> list[str]:
    output: list[str] = []
    for relationship_id in _image_ids(element):
        try:
            image = part.related_parts[relationship_id]
            if not image.content_type.startswith("image/"):
                raise ValueError("Embedded object is not an image")
            data = image.blob
            digest = hashlib.sha256(data).hexdigest()
            if digest not in cache:
                try:
                    cache[digest] = ocr_image(data)
                except ValueError as error:
                    cache[digest] = error
            recognized = cache[digest]
            if isinstance(recognized, ValueError):
                raise recognized
            if recognized:
                output.append(recognized)
                warnings.append(ParseWarning(
                    code="ocr_needs_review",
                    message="Text read from an embedded Word image needs source review, especially values and table cells",
                    location=location,
                ))
            else:
                warnings.append(ParseWarning(
                    code="image_no_text", message="No text was recovered from this image; a chart or diagram may still contain information",
                    location=location,
                ))
        except (KeyError, OcrUnavailable, ValueError) as error:
            warnings.append(ParseWarning(
                code="image_ocr_unavailable",
                message=f"Embedded image could not be read: {error}", location=location,
            ))
    return output


def _cell_text(cell, warnings: list[ParseWarning], location: str,
               cache: dict[str, str | ValueError], depth: int = 0) -> tuple[str, bool]:
    parts: list[str] = []
    needs_review = False
    for block in cell.iter_inner_content():
        if isinstance(block, Paragraph):
            value = normalize_inline(_format_runs(block))
            recognized = _image_text(block._p, block.part, location, warnings, cache)
            needs_review |= bool(recognized or fidelity_reason(value))
            if reason := fidelity_reason(value):
                warnings.append(ParseWarning("text_fidelity_needs_review", reason, location))
            if recognized:
                value = "\n".join([value, *recognized]).strip()
            if value:
                parts.append(value)
            if any(node.tag == qn("m:oMath") for node in block._p.iter()):
                needs_review = True
                warnings.append(ParseWarning("docx_equation_unavailable",
                    "Word equation layout has no verified text transcription", location))
                parts.append(f"Untranscribed equation at {location}; inspect the original source.")
            continue
        nested = block
        needs_review = True  # Nested cell relationships need review even without OCR.
        if depth >= 8:
            warnings.append(ParseWarning("nested_table_depth_limit", "Nested table depth exceeds eight", location))
            parts.append("Nested table content requires inspection in the original document.")
            continue
        for row_number, row in enumerate(nested.rows, start=1):
            values = [_cell_text(child, warnings, location, cache, depth + 1)[0] for child in row.cells]
            parts.append(f"Nested table row {row_number}: " + "; ".join(
                f"Column {index + 1}: {value}" for index, value in enumerate(values)))
    return normalize_multiline("\n".join(parts)), needs_review


def parse_docx_to_markdown(
    path: str | Path,
    source_name: str | None = None,
) -> ParsedDocument:
    input_path = Path(path)
    display_name = source_name or input_path.name
    document = Document(input_path)
    parts: list[str] = []
    warnings: list[ParseWarning] = []
    table_number = 0
    image_cache: dict[str, str | ValueError] = {}
    paragraph_number = 0

    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            paragraph_number += 1
            rendered = _paragraph_to_markdown(block)
            reason = fidelity_reason(block.text)
            if reason and rendered:
                warnings.append(ParseWarning("text_fidelity_needs_review", reason, f"paragraph {paragraph_number}"))
                rendered = f"<!-- review-required: text-fidelity -->\n\n{rendered}\n\n<!-- review-end -->"
            if rendered:
                parts.append(rendered)
            if any(node.tag == qn("m:oMath") for node in block._p.iter()):
                warnings.append(ParseWarning("docx_equation_unavailable",
                    "Word equation layout has no verified text transcription", f"paragraph {paragraph_number}"))
                parts.extend(("<!-- review-required: equation-unavailable -->",
                              f"Untranscribed equation at paragraph {paragraph_number}; inspect the original source.",
                              "<!-- review-end -->"))
            for recognized in _image_text(block._p, block.part, f"paragraph {paragraph_number}", warnings, image_cache):
                parts.extend(("<!-- review-required: docx-image-ocr -->", recognized, "<!-- review-end -->"))
            continue

        table_number += 1
        table_needs_review = False
        rows = []
        for row_number, row in enumerate(block.rows, start=1):
            cells = []
            for column_number, cell in enumerate(row.cells, start=1):
                value, needs_review = _cell_text(cell, warnings, f"table {table_number} row {row_number} column {column_number}", image_cache)
                cells.append(value)
                table_needs_review |= needs_review
            rows.append(cells)
        if not rows:
            warnings.append(
                ParseWarning(
                    code="empty_table",
                    message="An empty Word table was ignored",
                    location=f"table {table_number}",
                )
            )
            continue
        if any(cell.tables for row in block.rows for cell in row.cells):
            warnings.append(
                ParseWarning(
                    code="nested_table_review",
                    message="Nested table cells were retained as labeled text; verify their relationships",
                    location=f"table {table_number}",
                )
            )
        merged = any(len({id(cell._tc) for cell in row.cells}) != len(row.cells) for row in block.rows)
        if merged:
            table_needs_review = True
            warnings.append(ParseWarning("merged_table_review",
                "Merged Word cells repeat across grid positions; verify column relationships", f"table {table_number}"))
        first_data_row = 2
        if len(rows[0]) == 2 and rows[0][0].rstrip().endswith("?") and rows[0][1]:
            rows.insert(0, ["Question", "Answer"])
            first_data_row = 1
            warnings.append(
                ParseWarning(
                    code="table_header_inferred",
                    message="A two-column question/answer table had no visible header",
                    location=f"table {table_number}",
                )
            )
            table_needs_review = True
        table_markdown = (
            f"<!-- table: {table_number}; first_data_row: {first_data_row} -->\n"
            + render_markdown_table(rows)
        )
        if table_needs_review:
            parts.append("<!-- review-required: docx-table -->")
        parts.append(table_markdown)
        if table_needs_review:
            parts.append("<!-- review-end -->")

    # Check each header and footer once for embedded images
    visited = set()
    for section in document.sections:
        for label, region in (("header", section.header), ("footer", section.footer)):
            part = region.part
            if part.partname in visited:
                continue
            visited.add(part.partname)
            for recognized in _image_text(part.element, part, label, warnings, image_cache):
                parts.extend(("<!-- section-break -->", "<!-- review-required: docx-image-ocr -->",
                              f"Source document metadata:\n\n{label.title()} image text (OCR): {recognized}",
                              "<!-- review-end -->"))

    markdown = finish_markdown(parts)
    if "\ufffd" in markdown:
        warnings.append(ParseWarning(
            code="source_replacement_character",
            message=(
                "The DOCX contains Unicode replacement characters; the missing original "
                "characters cannot be reconstructed safely"
            ),
        ))
    if not markdown:
        warnings.append(
            ParseWarning(
                code="no_document_content",
                message="No readable paragraphs or tables were found in the DOCX file",
            )
        )

    return ParsedDocument(
        filename=display_name,
        document_type="docx",
        strategy_used="docx_structure_to_markdown",
        markdown=markdown,
        warnings=warnings,
    )
