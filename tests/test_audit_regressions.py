"""Evidence bugs found during the complete parsing-to-storage audit."""

from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch, Mock

from docx import Document
from docx.oxml import OxmlElement
import pymupdf as fitz

from gsme_rag.chunking import chunk_markdown
from gsme_rag.parsing import parse_to_markdown
from gsme_rag.parsing.fidelity import inspect_pdf_fidelity
from gsme_rag.parsing.markdown import render_markdown_table, split_table_row
from gsme_rag.storage.providers import CohereEmbeddings, FakeEmbeddings
from gsme_rag.storage.pipeline import search_evidence
from gsme_rag.storage.config import Settings
from gsme_rag.text import normalize_inline
from gsme_rag.version import PARSER_VERSION, CHUNKER_VERSION


class AuditRegressionTests(TestCase):
    def test_code_keeps_indentation_and_literal_page_markers(self):
        from gsme_rag.parsing.markdown import finish_markdown
        code = '```python\nif enabled:\n    print("<!-- page: 99 -->")\n\n\n    run()\n```'
        markdown = finish_markdown(["<!-- page: 1 -->", code, "Following prose."])
        chunks = chunk_markdown(markdown).chunks
        self.assertEqual(chunks[0].chunk_type, "code")
        self.assertIn(code, chunks[0].text)
        self.assertTrue(all(chunk.source_page_start == 1 for chunk in chunks))

    def test_generic_visual_gap_does_not_exclude_its_real_caption(self):
        text = ("<!-- page: 1 -->\n<!-- review-required: source-evidence -->\n\n"
                "Figure 1. Process diagram\n\n"
                "Visual evidence on page 1 requires source review; no verified description is available.\n"
                "<!-- review-end -->")
        chunks = chunk_markdown(text).chunks
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0].evidence_status, "pending")
        self.assertEqual(chunks[1].evidence_status, "excluded")

    def test_technical_symbols_are_not_compatibility_normalized(self):
        text = "x² + H₂O, 10 µm, 𝜎, ①"
        self.assertEqual(normalize_inline(text), text)

    def test_question_heading_pairs_keep_category_and_question(self):
        result = chunk_markdown("# General\n\n## What is Muse?\n\nMuse provides services.\n\n"
                                "## Where is support?\n\nSupport is global.")
        self.assertEqual(result.content_route, "faq")
        self.assertEqual([chunk.faq_question for chunk in result.chunks], ["What is Muse?", "Where is support?"])
        self.assertEqual(result.chunks[0].heading_path, ["General"])

    def test_faq_table_aliases_and_extra_columns_preserve_notes(self):
        markdown = render_markdown_table([["Notes", "Answers", "Questions"],
                                         ["Use the official portal.", "Global support.", "Where is support?"]])
        chunk = chunk_markdown(markdown).chunks[0]
        self.assertEqual(chunk.faq_question, "Where is support?")
        self.assertIn("Use the official portal.", chunk.text)
        styled = chunk_markdown("| **Question** | **Answer** |\n| --- | --- |\n| **What is Muse?** | Services. |").chunks[0]
        self.assertEqual(styled.faq_question, "What is Muse?")

    def test_native_numbered_sections_keep_chapter_and_nested_levels(self):
        from gsme_rag.parsing.pdf import _Line, _layout_markdown
        lines = [_Line(1, text, 70, 80 + index * 25, 450, 98 + index * 25, 11, 0)
                 for index, text in enumerate(["Chapter 1 Process", "1.1 Overview", "1.1.1 Detail", "Source evidence."])]
        result = chunk_markdown(_layout_markdown(lines)).chunks[0]
        self.assertEqual(result.heading_path, ["Chapter 1 Process", "1.1 Overview", "1.1.1 Detail"])

    def test_repeated_header_detection_counts_distinct_pages(self):
        from gsme_rag.parsing.pdf import _extract_lines
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((50, 25), "Visible twice on one page")
            page.insert_text((50, 45), "Visible twice on one page")
            doc.new_page().insert_text((50, 150), "Other page")
            self.assertEqual(sum(line.text == "Visible twice on one page" for line in _extract_lines(doc)), 2)

    def test_table_escape_roundtrip_and_ragged_oversized_row(self):
        row = [r"path\|file", "A | B", "last|", r"C:\new\file"]
        rendered = render_markdown_table([row, row])
        self.assertEqual(split_table_row(rendered.splitlines()[0]), row)
        malformed = "| Name | Value |\n| --- | --- |\n| A | 10 | " + "extra " * 100 + "|"
        chunks = chunk_markdown(malformed, max_chars=200).chunks
        self.assertTrue(all(chunk.review_required for chunk in chunks))
        self.assertTrue(any("Column 3" in chunk.text for chunk in chunks))
        self.assertTrue(all(len(chunk.text) <= 200 for chunk in chunks))

    def test_docx_hyperlink_nested_table_and_damaged_text_are_retained(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "nested.docx"
            doc = Document()
            p = doc.add_paragraph("See ")
            link = OxmlElement("w:hyperlink")
            run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "the support portal"
            run.append(text); link.append(run); p._p.append(link)
            doc.add_paragraph("Damaged \ufffd source")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Field"; table.cell(0, 1).text = "Value"
            cells = table.add_row().cells; cells[0].text = "Nested evidence"
            nested = cells[1].add_table(rows=1, cols=2)
            nested.cell(0, 0).text = "LOT A17"; nested.cell(0, 1).text = "92%"
            doc.save(source)
            parsed = parse_to_markdown(source)
            self.assertIn("the support portal", parsed.markdown)
            self.assertIn("LOT A17", parsed.markdown)
            chunks = chunk_markdown(parsed.markdown, parsed.filename, metadata=parsed.metadata).chunks
            self.assertTrue(all(c.review_required for c in chunks if "\ufffd" in c.text or "LOT A17" in c.text))
            self.assertEqual(parsed.metadata.parser_version, PARSER_VERSION)
            self.assertTrue(all(c.chunker_version == CHUNKER_VERSION for c in chunks))

    def test_simple_pdf_formula_path_does_not_bypass_fidelity(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "simple.pdf"
            with fitz.open() as doc:
                page = doc.new_page(); page.insert_text((70, 100), "DPPM = (")
                doc.save(source)
            with patch.dict("os.environ", {"GSME_PDF_BACKEND": "native"}):
                parsed = parse_to_markdown(source)
            self.assertTrue(parsed.evidence_diagnostics)
            chunks = chunk_markdown(parsed.markdown, source.name, metadata=parsed.metadata).chunks
            self.assertTrue(all(c.review_required for c in chunks))

    def test_fidelity_ocr_zero_budget_is_strict(self):
        with fitz.open() as doc:
            page = doc.new_page(); page.insert_text((70, 100), "Yield = (")
            with patch("gsme_rag.parsing.fidelity.ocr_image") as ocr:
                inspect_pdf_fidelity(doc, max_ocr_regions=0)
                ocr.assert_not_called()

    def test_source_modified_during_parsing_cannot_get_a_valid_source_hash(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "faq.docx"
            doc = Document(); doc.add_paragraph("Original"); doc.save(source)
            from gsme_rag.parsing.docx import parse_docx_to_markdown
            def modified(*args, **kwargs):
                parsed = parse_docx_to_markdown(*args, **kwargs)
                source.write_bytes(b"changed bytes")
                return parsed
            with patch("gsme_rag.parsing.parser.parse_docx_to_markdown", side_effect=modified):
                with self.assertRaisesRegex(ValueError, "changed during"):
                    parse_to_markdown(source)

    def test_empty_index_does_not_send_query_to_external_provider(self):
        store = Mock(); store.has_indexed_evidence.return_value = False
        provider = Mock(model_version="m"); vectors = Mock(collection="c")
        self.assertEqual(search_evidence(store, vectors, provider, Settings("unused", "unused", Path("data")),
                                         "t", "question"), [])
        provider.encode_tokens.assert_not_called(); provider.embed_query.assert_not_called()

    def test_provider_rejects_nan_and_zero_vectors(self):
        provider = CohereEmbeddings.__new__(CohereEmbeddings); provider.dimension = 2
        for values in ([float("nan"), 1.0], [0.0, 0.0]):
            provider.client = Mock()
            provider.client.embed.return_value = SimpleNamespace(embeddings=SimpleNamespace(float=[values]))
            with self.assertRaisesRegex(ValueError, "non-finite"):
                provider.embed_documents(["approved synthetic source"])

    def test_fake_tokenizer_treats_special_token_text_as_source_text(self):
        self.assertTrue(FakeEmbeddings(8).encode_tokens("Literal <|endoftext|> in documentation"))

    def test_docling_image_child_heading_does_not_leak_to_next_page(self):
        from gsme_rag.parsing.docling_pdf import _normalize_document
        def item(label, page, text=""):
            return SimpleNamespace(label=label, text=text, prov=[SimpleNamespace(page_no=page)],
                                   caption_text=lambda _: "Figure 1. Source image")
        class Source:
            def iterate_items(self):
                yield item("picture", 1), 0
                yield item("section_header", 1, "LOT A17"), 1
                yield item("text", 2, "Selectable next page text."), 2
        parsed = _normalize_document(Source(), "source.pdf", set())
        self.assertNotIn("## LOT A17", parsed.markdown)
        chunks = chunk_markdown(parsed.markdown).chunks
        self.assertTrue(all(not chunk.review_required for chunk in chunks if chunk.source_page_start == 2))

    def test_unlocated_dropped_cells_keep_raw_evidence(self):
        from gsme_rag.parsing.docling_pdf import _cell_loss_evidence
        with fitz.open() as source:
            source.new_page().insert_text((50, 100), "Unrelated page")
            evidence, warnings, pages = _cell_loss_evidence(source, [{
                "bbox": [1, 2, 3, 4], "missing_cells": [{"id": 12, "text": "Lost source cell", "bbox": [1, 2, 3, 4]}]
            }])
        self.assertEqual(evidence[0].original_text, "Lost source cell")
        self.assertIsNone(evidence[0].page)
        self.assertTrue(warnings)

    def test_source_evaluation_rejects_unpinned_cases_before_external_calls(self):
        import json
        from gsme_rag.storage.evaluate import evaluate_cases
        with TemporaryDirectory() as directory:
            case_file = Path(directory) / "cases.json"
            case_file.write_text(json.dumps({"source": "source_reviewed", "cases": [
                {"tenant_id": "t", "question": "What is Muse?", "expected_chunk_ids": ["c"]}]}), encoding="utf-8")
            provider = Mock()
            with self.assertRaisesRegex(ValueError, "pin"):
                evaluate_cases(Mock(), Mock(), provider, Mock(), case_file)
            provider.embed_query.assert_not_called()
