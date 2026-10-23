"""OCR in Word, standalone scans, and mixed digital/image PDFs."""

from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
import pymupdf

from gsme_rag.chunking import chunk_markdown, chunk_faqs, chunk_papers
from gsme_rag.models import ParsedDocument
from gsme_rag.parsing import parse_to_markdown
from gsme_rag.parsing.docling_pdf import _normalize_document
from gsme_rag.parsing.pdf import _Line, _insert_image_ocr_at_figure, _supplement_image_ocr


def sample_image() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=720, height=150)
    page.insert_text((25, 65), "Picture pitch 0.127 mm", fontsize=24)
    result = page.get_pixmap().tobytes("png")
    document.close()
    return result


class ImageOcrTests(unittest.TestCase):
    def test_image_ocr_is_inserted_after_figure_before_later_heading(self):
        markdown = (
            "<!-- page: 2 -->\n\nFigure 1. Raster evidence\n\n"
            "## 2.1 Expected behavior\n\nLater prose."
        )
        result = _insert_image_ocr_at_figure(markdown, "LOT A17")
        self.assertLess(result.index("Image text (OCR): LOT A17"), result.index("## 2.1"))
        chunks = chunk_markdown(result, "sample.pdf").chunks
        ocr = next(chunk for chunk in chunks if "LOT A17" in chunk.text)
        self.assertNotIn("2.1 Expected behavior", ocr.heading_path)
        self.assertTrue(ocr.review_required)

    def test_numeric_visual_page_keeps_formula_and_chart_review_state(self):
        def item(label, page, text="", caption=""):
            return SimpleNamespace(
                label=label, prov=[SimpleNamespace(page_no=page)], text=text,
                caption_text=lambda doc: caption,
            )

        class Document:
            def iterate_items(self):
                yield item("section_header", 3, "Charts and formulas need review"), 1
                yield item("section_header", 3, "3.1 Formula source line"), 1
                yield item("text", 3, "GDPW = 3.14 x wafer radius"), 1
                yield item("picture", 3, caption="Figure 2. Synthetic chart"), 1

        parsed = _normalize_document(Document(), "sample.pdf", {3}, {3})
        result = chunk_markdown(parsed.markdown, parsed.filename)
        self.assertTrue(all(chunk.review_required for chunk in result.chunks if chunk.source_page_start == 3))

    def test_faq_and_paper_have_separate_automatic_routes(self):
        faq = "# General\n\n**What is the pitch?**\n\nIt is 0.127 mm."
        paper = "# Measurements\n\n**Why is the pitch different?**\n\nA paper studies its variation."
        routed = chunk_markdown(faq)
        self.assertEqual(routed.content_route, "faq")
        self.assertEqual(routed.chunks[0].chunk_type, "question_answer")
        forced = chunk_papers(paper)
        self.assertEqual(forced.content_route, "paper")
        self.assertFalse(any(c.chunk_type == "question_answer" for c in forced.chunks))
        self.assertEqual(chunk_faqs(faq).content_route, "faq")

    def test_docx_question_image_answer_and_table_cell_are_review_marked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.docx"
            document = Document()
            question = document.add_paragraph()
            question.add_run("What is the pitch?").bold = True
            document.add_paragraph().add_run().add_picture(BytesIO(sample_image()))
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Question"
            table.cell(0, 1).text = "Answer"
            table.cell(1, 0).text = "What is the pitch in the table?"
            table.cell(1, 1).paragraphs[0].add_run().add_picture(BytesIO(sample_image()))
            document.save(path)
            with patch("gsme_rag.parsing.docx.ocr_image", return_value="0.127 mm") as ocr:
                parsed = parse_to_markdown(path)
            self.assertEqual(ocr.call_count, 1)  # duplicate bytes OCRed once
            self.assertEqual([w.code for w in parsed.warnings].count("ocr_needs_review"), 2)
            chunks = chunk_markdown(parsed.markdown, path.name, metadata=parsed.metadata).chunks
            self.assertEqual(len([c for c in chunks if c.chunk_type == "question_answer"]), 1)
            self.assertTrue(any(c.review_required and "0.127 mm" in c.text for c in chunks))
            self.assertTrue(next(c for c in chunks if c.chunk_type == "table_row").review_required)

    def test_header_image_does_not_become_last_faq_answer(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "header.docx"
            document = Document()
            document.add_paragraph().add_run("What is the pitch?").bold = True
            document.add_paragraph("The answer is in the body.")
            document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(sample_image()))
            document.save(path)
            with patch("gsme_rag.parsing.docx.ocr_image", return_value="Confidential header text"):
                parsed = parse_to_markdown(path)
            chunks = chunk_markdown(parsed.markdown, path.name, metadata=parsed.metadata).chunks
            answer = next(c for c in chunks if c.chunk_type == "question_answer")
            self.assertNotIn("Confidential header", answer.text)
            self.assertTrue(any(c.review_required and "Confidential header" in c.text for c in chunks))

    def test_standalone_image_is_review_marked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scan.png"
            path.write_bytes(sample_image())
            with patch("gsme_rag.parsing.image.ocr_image", return_value="Pitch 0.127 mm"):
                parsed = parse_to_markdown(path)
            self.assertEqual(parsed.document_type, "image")
            chunks = chunk_markdown(parsed.markdown, path.name, metadata=parsed.metadata).chunks
            self.assertEqual(len(chunks), 1)
            self.assertTrue(chunks[0].review_required)

    def test_pdf_selectable_text_and_raster_text_keep_separate_review_states(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((50, 80), "This document contains ordinary selectable manufacturing process text.")
            page.insert_text((50, 110), "The illustration below contains a separate measurement on a raster image.")
            page.insert_image(pymupdf.Rect(50, 220, 510, 340), stream=sample_image())
            document.save(path)
            document.close()
            with patch.dict("os.environ", {"GSME_PDF_BACKEND": "native"}), patch(
                "gsme_rag.parsing.pdf.ocr_image", return_value="Picture pitch 0.127 mm"
            ):
                parsed = parse_to_markdown(path)
            chunks = chunk_markdown(parsed.markdown, path.name, metadata=parsed.metadata).chunks
            self.assertTrue(any(c.review_required and "0.127" in c.text for c in chunks))
            self.assertTrue(any(not c.review_required and "ordinary selectable" in c.text for c in chunks))

    def test_docling_layout_supplements_unread_raster_image(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(pymupdf.Rect(50, 220, 510, 340), stream=sample_image())
        parsed = ParsedDocument("scan.pdf", "pdf", "docling_layout_ocr_to_markdown",
                                "<!-- page: 1 -->\n\nSome selectable prose.\n")
        with patch("gsme_rag.parsing.pdf.ocr_image", return_value="Picture pitch 0.127 mm"):
            result = _supplement_image_ocr(document, parsed, [])
        self.assertIn("Picture pitch 0.127 mm", result.markdown)
        self.assertTrue(any(c.review_required for c in chunk_markdown(result.markdown, "scan.pdf").chunks))
        document.close()

    def test_docling_existing_raster_text_is_reviewed_not_promoted_to_heading(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(pymupdf.Rect(50, 220, 510, 340), stream=sample_image())
        parsed = ParsedDocument(
            "scan.pdf", "pdf", "docling_layout_ocr_to_markdown",
            "<!-- page: 1 -->\n\n## Picture pitch 0.127 mm\n\n"
            "<!-- page: 2 -->\n\nVerified next-page prose.\n",
        )
        with patch("gsme_rag.parsing.pdf.ocr_image", return_value="Picture pitch 0.127 mm"):
            result = _supplement_image_ocr(document, parsed, [])
        self.assertNotIn("## Picture pitch", result.markdown)
        self.assertIn("Image text (OCR): Picture pitch 0.127 mm", result.markdown)
        self.assertIn("visual_evidence_missing", [warning.code for warning in result.warnings])
        chunk = next(c for c in chunk_markdown(result.markdown, "scan.pdf").chunks if "0.127" in c.text)
        self.assertTrue(chunk.review_required)
        next_page = next(c for c in chunk_markdown(result.markdown, "scan.pdf").chunks
                         if "next-page" in c.text)
        self.assertEqual(next_page.source_page_start, 2)
        self.assertFalse(next_page.review_required)
        document.close()

    def test_conflicting_docling_and_supplemental_ocr_are_one_review_unit(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(pymupdf.Rect(50, 220, 510, 340), stream=sample_image())
        docling_reading = "OCR TEST: YIELD 92 PERCENT LOT A17 / REVIEW NUMBERS"
        supplemental_reading = "OCR TEST: YIELD 92 PERCENT LOT Al7 / REVIEW NUMBERS"
        parsed = ParsedDocument(
            "scan.pdf", "pdf", "docling_layout_ocr_to_markdown",
            f"<!-- page: 1 -->\n\n## {docling_reading}\n\nFigure caption.\n",
        )
        with patch("gsme_rag.parsing.pdf.ocr_image", return_value=supplemental_reading):
            result = _supplement_image_ocr(document, parsed, [])
        self.assertNotIn(f"## {docling_reading}", result.markdown)
        self.assertEqual(result.markdown.count(docling_reading), 1)
        self.assertEqual(result.markdown.count(supplemental_reading), 1)
        self.assertIn("ocr_reading_conflict", [warning.code for warning in result.warnings])
        matching = [c for c in chunk_markdown(result.markdown, "scan.pdf").chunks
                    if "LOT A17" in c.text or "LOT Al7" in c.text]
        self.assertEqual(len(matching), 1)
        self.assertTrue(matching[0].review_required)
        self.assertIn("Docling OCR:", matching[0].text)
        self.assertIn("Supplemental OCR:", matching[0].text)
        document.close()

    def test_docling_restores_numbered_section_context_from_selectable_text(self):
        document = pymupdf.open()
        document.new_page()
        parsed = ParsedDocument(
            "manual.pdf", "pdf", "docling_layout_ocr_to_markdown",
            "<!-- page: 1 -->\n\n# Chapter 1\n\nOverview of the flow\n\n"
            "A question in a technical manual\n\nWhat is the purpose?\n",
        )
        lines = [
            _Line(1, "1.1 Overview of the flow", 50, 80, 240, 95, 12, 0),
            _Line(1, "1.3 A question in a technical manual", 50, 160, 300, 175, 12, 0),
        ]
        result = _supplement_image_ocr(document, parsed, lines)
        self.assertIn("## 1.1 Overview of the flow", result.markdown)
        self.assertIn("## 1.3 A question in a technical manual", result.markdown)
        chunks = chunk_markdown(result.markdown, result.filename).chunks
        question = next(c for c in chunks if "What is the purpose?" in c.text)
        self.assertEqual(question.heading_path, ["Chapter 1", "1.3 A question in a technical manual"])
        document.close()

    def test_docling_recovers_omitted_numeric_selectable_labels_for_review(self):
        document = pymupdf.open()
        document.new_page()
        parsed = ParsedDocument(
            "chart.pdf", "pdf", "docling_layout_ocr_to_markdown",
            "<!-- page: 1 -->\n\n## 3.2 Chart label mapping\n\nFigure caption.\n\n## 3.3 Notes\n\nClosing text.\n",
        )
        lines = [
            _Line(1, "A: 20 percent", 100, 200, 180, 215, 11, 0),
            _Line(1, "B: 80 percent", 100, 220, 180, 235, 11, 0),
            _Line(1, "3.3 Notes", 50, 300, 150, 315, 14, 1),
        ]
        result = _supplement_image_ocr(document, parsed, lines)
        self.assertIn("A: 20 percent", result.markdown)
        self.assertIn("B: 80 percent", result.markdown)
        self.assertIn("docling_text_omission", [warning.code for warning in result.warnings])
        recovered = next(c for c in chunk_markdown(result.markdown, "chart.pdf").chunks if "A: 20" in c.text)
        self.assertEqual(recovered.heading_path, ["3.2 Chart label mapping"])
        self.assertTrue(recovered.review_required)
        document.close()


if __name__ == "__main__":
    unittest.main()
