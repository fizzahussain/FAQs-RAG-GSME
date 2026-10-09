import csv
from pathlib import Path
import tempfile
import unittest

from docx import Document
import pymupdf as fitz

from gsme_rag.chunking import chunk_markdown, compare_strategies
from gsme_rag.parsing import parse_to_markdown
from gsme_rag.parsing.schedules import filter_schedule_records, parse_schedule_records
from gsme_rag.parsing.schedules import resolve_schedule_kind


class ParsingStageTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_docx_paragraphs_become_markdown_without_chunks(self) -> None:
        path = self.root / "faqs.docx"
        document = Document()
        document.add_heading("General", level=1)
        question = document.add_paragraph()
        question.add_run("01 What is Muse?").bold = True
        document.add_paragraph("Muse provides semiconductor services.")
        document.save(path)

        result = parse_to_markdown(path)

        self.assertEqual(result.document_type, "docx")
        self.assertEqual(result.strategy_used, "docx_structure_to_markdown")
        self.assertIn("# General", result.markdown)
        self.assertIn("**01 What is Muse?**", result.markdown)
        self.assertNotIn("chunk_id", result.markdown)

    def test_docx_table_becomes_markdown_table(self) -> None:
        path = self.root / "faqs.docx"
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Question"
        table.rows[0].cells[1].text = "Answer"
        row = table.add_row().cells
        row[0].text = "What is Muse?"
        row[1].text = "Muse provides semiconductor services."
        document.save(path)

        result = parse_to_markdown(path)

        self.assertIn("| Question | Answer |", result.markdown)
        self.assertIn("| What is Muse? |", result.markdown)

    def test_docx_headerless_faq_table_keeps_first_question(self) -> None:
        path = self.root / "headerless.docx"
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "What is Muse?"
        table.rows[0].cells[1].text = "Muse provides services."
        document.save(path)

        result = parse_to_markdown(path)
        chunks = chunk_markdown(result.markdown, path.name)

        self.assertIn("| Question | Answer |", result.markdown)
        self.assertEqual(chunks.chunks[0].table_row_number, 1)
        self.assertIn("What is Muse?", chunks.chunks[0].text)
        self.assertEqual(result.warnings[0].code, "table_header_inferred")

    def test_csv_becomes_markdown_table_without_schedule_semantics(self) -> None:
        path = self.root / "shared_schedules.csv"
        with path.open("w", newline="", encoding="utf-8") as target:
            writer = csv.writer(target)
            writer.writerow(
                ["Tech (nm)", "Flavor", "Trial", "Final", "Tapeout", "Est. Ship"]
            )
            writer.writerow(
                [28, "HPC+ RF", "08/26/2026", "09/02/2026", "09/09/2026", "12/17/2026"]
            )

        result = parse_to_markdown(path)

        self.assertEqual(result.strategy_used, "csv_to_markdown_table")
        self.assertIn("# Shared Schedules", result.markdown)
        self.assertIn("| Tech (nm) | Flavor |", result.markdown)
        self.assertNotIn("schedule_id", result.markdown)

    def test_csv_keeps_extra_cells_instead_of_truncating(self) -> None:
        path = self.root / "uneven.csv"
        path.write_text("Technology,Flavor\n28,HPC+ RF,Extra detail\n", encoding="utf-8")

        result = parse_to_markdown(path)

        self.assertIn("| Technology | Flavor | Column 3 |", result.markdown)
        self.assertIn("| 28 | HPC+ RF | Extra detail |", result.markdown)
        self.assertEqual(result.warnings[0].code, "uneven_csv_row")

    def test_schedule_records_are_typed_and_keep_repeated_technology(self) -> None:
        path = self.root / "full_schedules.csv"
        path.write_text(
            "technology_nm,flavor,trial_deadline,final_deadline,tapeout_date,estimated_ship_date\n"
            "28,HPC+ RF,08/26/2026,09/02/2026,09/09/2026,12/17/2026\n"
            "28,HPC+ RF,10/28/2026,11/04/2026,11/11/2026,02/18/2027\n",
            encoding="utf-8",
        )
        records = parse_schedule_records(path)
        self.assertEqual([row["source_row_number"] for row in records], [2, 3])
        self.assertEqual(records[0]["technology_nm"], 28)
        self.assertEqual(records[0]["tapeout_date"].isoformat(), "2026-09-09")
        matched = filter_schedule_records(records, 28, "hpc+ rf", "2026-10-01")
        self.assertEqual([row["source_row_number"] for row in matched], [3])

    def test_schedule_record_invalid_date_is_reported_with_row(self) -> None:
        path = self.root / "schedule.csv"
        path.write_text(
            "technology_nm,flavor,trial_deadline,final_deadline,tapeout_date,estimated_ship_date\n"
            "28,LP,08/26/2026,09/02/2026,not-a-date,12/17/2026\n",
            encoding="utf-8",
        )
        with self.assertRaisesRegex(ValueError, "row 2: tapeout_date"):
            parse_schedule_records(path)

    def test_generically_named_schedule_requires_explicit_kind_for_context(self) -> None:
        path = self.root / "no name in file.csv"
        path.write_text(
            "technology_nm,flavor,trial_deadline,final_deadline,tapeout_date,estimated_ship_date\n"
            "28,HPC+ RF,8/26/2026,9/2/2026,9/9/2026,12/17/2026\n",
            encoding="utf-8",
        )
        unknown = parse_to_markdown(path)
        labeled = parse_to_markdown(path, schedule_kind="shared")
        self.assertEqual(unknown.warnings[0].code, "schedule_kind_unknown")
        self.assertTrue(labeled.markdown.startswith("# Shared Schedules"))
        self.assertEqual(labeled.filename, "no name in file.csv")
        self.assertEqual(labeled.warnings, [])
        self.assertEqual(unknown.metadata.classification_status, "needs_schedule_kind")
        self.assertEqual(labeled.metadata.schedule_kind_source, "user")
        with self.assertRaisesRegex(ValueError, "Schedule type is unknown"):
            chunk_markdown(unknown.markdown, path.name, metadata=unknown.metadata)
        structured_only = chunk_markdown(labeled.markdown, path.name, metadata=labeled.metadata)
        self.assertEqual(structured_only.chunks, [])
        self.assertEqual(structured_only.warnings[0].code, "structured_records_only")
        discovery = chunk_markdown(
            labeled.markdown, path.name, metadata=labeled.metadata,
            semantic_schedule_rows=True,
        )
        self.assertEqual(len(discovery.chunks), 1)
        self.assertEqual(discovery.chunks[0].schedule_kind, "shared")
        self.assertEqual(discovery.chunks[0].source_row_number, 2)
        self.assertEqual(discovery.chunks[0].source_id, labeled.metadata.source_id)
        with self.assertRaisesRegex(ValueError, "row-preserving"):
            chunk_markdown(
                labeled.markdown, path.name, strategy="fixed_size",
                metadata=labeled.metadata, semantic_schedule_rows=True,
            )
        with self.assertRaisesRegex(ValueError, "filename says full schedules"):
            resolve_schedule_kind("full_schedules.csv", "shared")

    def test_adjacent_italic_runs_are_one_markdown_span(self) -> None:
        path = self.root / "story.docx"
        document = Document()
        paragraph = document.add_paragraph("I wondered, ")
        paragraph.add_run("Why compare ").italic = True
        paragraph.add_run("our timelines?").italic = True
        document.save(path)
        self.assertIn("*Why compare our timelines?*", parse_to_markdown(path).markdown)

    def test_pdf_reading_order_becomes_markdown(self) -> None:
        path = self.root / "faqs.pdf"
        document = fitz.open()
        page = document.new_page(width=612, height=792)
        page.insert_text((50, 70), "General", fontname="hebo", fontsize=15)
        page.insert_text((50, 105), "01 What is Muse?", fontname="hebo", fontsize=11)
        page.insert_text(
            (50, 130),
            "Muse provides semiconductor services.",
            fontname="helv",
            fontsize=11,
        )
        document.save(path)
        document.close()

        result = parse_to_markdown(path)

        self.assertEqual(result.document_type, "pdf")
        self.assertIn("<!-- page: 1 -->", result.markdown)
        self.assertIn("**01 What is Muse?**", result.markdown)

    def test_pdf_equal_size_category_headings_are_siblings(self) -> None:
        path = self.root / "faqs.pdf"
        document = fitz.open()
        for category, question, answer in [
            ("General", "01 Who do I contact?", "Email support."),
            ("TapeOut Services", "01 What is the minimum area?", "It depends."),
        ]:
            page = document.new_page(width=612, height=792)
            page.insert_text((50, 70), category, fontname="hebo", fontsize=15)
            page.insert_text((50, 105), question, fontname="hebo", fontsize=11)
            page.insert_text((50, 130), answer, fontname="helv", fontsize=11)
        document.save(path)
        document.close()

        parsed = parse_to_markdown(path)
        self.assertIn("# General", parsed.markdown)
        self.assertIn("# TapeOut Services", parsed.markdown)
        self.assertNotIn("## TapeOut Services", parsed.markdown)
        chunks = chunk_markdown(parsed.markdown, path.name, metadata=parsed.metadata).chunks
        self.assertEqual([c.category for c in chunks], ["General", "TapeOut Services"])
        self.assertEqual(chunks[1].heading_path, ["TapeOut Services"])

    def test_pdf_without_selectable_text_returns_warning(self) -> None:
        path = self.root / "empty.pdf"
        document = fitz.open()
        document.new_page(width=612, height=792)
        document.save(path)
        document.close()

        result = parse_to_markdown(path)

        self.assertEqual(result.markdown, "")
        self.assertEqual(result.warnings[0].code, "no_pdf_text")


class ChunkingStageTests(unittest.TestCase):
    def test_pdf_sibling_categories_replace_general_in_chunks(self) -> None:
        markdown = (
            "<!-- page: 1 -->\n\n## General\n\n**01 Who do I contact?**\n\nEmail support.\n\n"
            "<!-- page: 2 -->\n\n## TapeOut Services\n\n**01 What is the minimum area?**\n\n"
            "It depends on technology.\n\n## Additional Services\n\n"
            "**01 Can I order more samples?**\n\nYes."
        )
        chunks = chunk_markdown(markdown, "faq.pdf").chunks
        self.assertEqual(
            [chunk.category for chunk in chunks],
            ["General", "TapeOut Services", "Additional Services"],
        )
        self.assertEqual(chunks[1].heading_path, ["TapeOut Services"])
        self.assertTrue(chunks[1].text.startswith("# TapeOut Services"))
        self.assertNotIn("General", chunks[1].text)

    def test_wrapped_bold_pdf_question_is_one_complete_pair(self) -> None:
        markdown = (
            "## TapeOut Services\n\n"
            "**09 Which technologies are supported?**\n\nAll TSMC technologies.\n\n"
            "**10 Can I access the standard cells and I/O libraries**\n\n"
            "**through Muse Semiconductor?**\n\nYes, request access by email."
        )
        chunks = chunk_markdown(markdown, "faq.pdf").chunks
        self.assertEqual(len(chunks), 2)
        self.assertNotIn("10 Can I access", chunks[0].text)
        self.assertIn(
            "**10 Can I access the standard cells and I/O libraries through Muse Semiconductor?**",
            chunks[1].text,
        )
        self.assertIn("Yes, request access by email.", chunks[1].text)

    def test_pdf_table_category_column_sets_explicit_category(self) -> None:
        markdown = (
            "# FAQ Table\n\n"
            "| Category | Question | Answer | Page |\n| --- | --- | --- | --- |\n"
            "| General | Who do I contact? | Email support. | 1 |\n"
            "| TapeOut Services | What is the minimum area? | It depends. | 2 |"
        )
        chunks = chunk_markdown(markdown, "faq.pdf").chunks
        self.assertEqual([chunk.category for chunk in chunks], ["General", "TapeOut Services"])
        self.assertEqual(chunks[1].heading_path, ["FAQ Table"])
        self.assertEqual(chunks[1].source_page_start, 2)

    def test_auto_routes_mixed_document_by_block(self) -> None:
        markdown = (
            "# Guide\n\n**What is the deadline?**\n\nThe final review is Friday.\n\n"
            "| Question | Answer |\n| --- | --- |\n| Who approves? | The lead. |\n\n"
            "A separate narrative paragraph.\n\nAnother independent observation."
        )
        result = chunk_markdown(markdown)
        self.assertEqual(
            [chunk.chunk_type for chunk in result.chunks],
            ["question_answer", "table_row", "section", "section"],
        )
        self.assertTrue(all(chunk.strategy_used == "auto" for chunk in result.chunks))
        self.assertEqual(len({chunk.source_id for chunk in result.chunks}), 1)
        self.assertEqual(
            [chunk.chunk_id for chunk in result.chunks],
            [chunk.chunk_id for chunk in chunk_markdown(markdown).chunks],
        )

    def test_edited_markdown_changes_hash_and_warns(self) -> None:
        from gsme_rag.models import DocumentMetadata
        metadata = DocumentMetadata(
            source_id="doc-abc", content_sha256="abc", normalized_sha256="old",
            source_file="note.docx", document_type="docx", document_class="document",
        )
        result = chunk_markdown("A revised observation.", source_file="note.docx", metadata=metadata)
        self.assertEqual(result.warnings[0].code, "markdown_edited")
        self.assertEqual(result.chunks[0].source_id, "doc-abc")

    def test_question_and_answer_stay_in_one_chunk(self) -> None:
        markdown = """# General

**01 What is Muse?**

Muse provides semiconductor services.
"""

        result = chunk_markdown(markdown, "faqs.md")

        self.assertEqual(len(result.chunks), 1)
        self.assertEqual(result.chunks[0].chunk_type, "question_answer")
        self.assertIn("What is Muse?", result.chunks[0].text)
        self.assertIn("semiconductor services", result.chunks[0].text)
        self.assertEqual(result.chunks[0].heading_path, ["General"])

    def test_markdown_table_creates_one_chunk_per_data_row(self) -> None:
        markdown = """# Shared Schedules

| Tech | Flavor |
| --- | --- |
| 28 | HPC+ RF |
| 65 | MS RF LP |
"""

        result = chunk_markdown(markdown, "shared_schedules.md")

        self.assertEqual(len(result.chunks), 2)
        self.assertTrue(all(chunk.chunk_type == "table_row" for chunk in result.chunks))
        self.assertTrue(all("| Tech | Flavor |" in chunk.text for chunk in result.chunks))

    def test_long_table_row_is_not_split_from_its_header(self) -> None:
        markdown = "# FAQ\n\n| Question | Answer |\n| --- | --- |\n"
        markdown += f"| What happened? | {'Long answer ' * 40}|\n"
        result = chunk_markdown(markdown, max_chars=200)

        self.assertEqual(len(result.chunks), 1)
        self.assertIn("| Question | Answer |", result.chunks[0].text)
        self.assertIn("What happened?", result.chunks[0].text)
        self.assertEqual(result.warnings[0].code, "oversized_unit")

    def test_table_page_and_row_number_are_preserved(self) -> None:
        markdown = "# FAQ\n\n| Question | Answer | Page |\n| --- | --- | --- |\n| Why? | Because. | 3 |"
        result = chunk_markdown(markdown)

        self.assertEqual(result.chunks[0].source_page_start, 3)
        self.assertEqual(result.chunks[0].table_row_number, 2)

    def test_pdf_table_row_page_range_is_preserved(self) -> None:
        markdown = "# FAQ\n\n| Question | Answer | Page |\n| --- | --- | --- |\n| Why? | A long answer. | 3-4 |"
        result = chunk_markdown(markdown)
        self.assertEqual(result.chunks[0].source_page_start, 3)
        self.assertEqual(result.chunks[0].source_page_end, 4)

    def test_generic_pdf_prose_keeps_page_boundaries(self) -> None:
        markdown = "<!-- page: 1 -->\n\n# History\n\nA first-page fact.\n\n<!-- page: 2 -->\n\nA second-page fact."
        result = chunk_markdown(markdown)

        self.assertEqual(len(result.chunks), 2)
        self.assertEqual([chunk.source_page_start for chunk in result.chunks], [1, 2])
        self.assertTrue(all(chunk.heading_path == ["History"] for chunk in result.chunks))

    def test_long_faq_parts_share_unit_id(self) -> None:
        markdown = "# General\n\n**01 Why?**\n\n" + "This answer is detailed. " * 30
        result = chunk_markdown(markdown, max_chars=220)

        self.assertGreater(len(result.chunks), 1)
        self.assertEqual(len({chunk.unit_id for chunk in result.chunks}), 1)
        self.assertEqual(
            [chunk.part_index for chunk in result.chunks],
            list(range(len(result.chunks))),
        )
        self.assertTrue(all(chunk.part_count == len(result.chunks) for chunk in result.chunks))

    def test_fixed_baseline_and_hybrid_support_narrative(self) -> None:
        markdown = "# History\n\n" + "The team worked on a new service. " * 35
        fixed = chunk_markdown(markdown, max_chars=250, strategy="fixed_size")
        hybrid = chunk_markdown(
            markdown, max_chars=250, strategy="hybrid", overlap_chars=40
        )

        self.assertGreater(len(fixed.chunks), 1)
        self.assertTrue(all(chunk.chunk_type == "fixed_size" for chunk in fixed.chunks))
        self.assertGreater(len(hybrid.chunks), 1)
        self.assertTrue(all(chunk.chunk_type == "section" for chunk in hybrid.chunks))
        self.assertTrue(all(chunk.character_count <= 250 for chunk in hybrid.chunks))

    def test_paragraph_aware_keeps_narrative_paragraphs_separate(self) -> None:
        markdown = "First topic has its own evidence.\n\nSecond topic has other evidence."
        result = chunk_markdown(markdown, strategy="paragraph_aware")
        self.assertEqual(len(result.chunks), 2)
        self.assertEqual([c.chunk_type for c in result.chunks], ["section", "section"])
        self.assertIn("First topic", result.chunks[0].text)
        self.assertNotIn("Second topic", result.chunks[0].text)

    def test_hybrid_overlap_uses_full_sentence(self) -> None:
        markdown = "First sentence ends clearly. " * 10
        result = chunk_markdown(markdown, strategy="hybrid", max_chars=210, overlap_chars=50)
        self.assertGreater(len(result.chunks), 1)
        self.assertTrue(all(c.character_count <= 210 for c in result.chunks))
        self.assertTrue(result.chunks[1].text.startswith("First sentence ends clearly."))

    def test_hybrid_does_not_repeat_unbalanced_emphasis(self) -> None:
        from gsme_rag.chunking.chunker import _overlap_parts

        parts = ["A long reflection. *Why this? And what next?*", "A new paragraph begins."]
        result = _overlap_parts(parts, 200, 80)
        self.assertEqual(result[1], "A new paragraph begins.")

    def test_unknown_strategy_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            chunk_markdown("Some text", strategy="unknown")

    def test_comparison_excludes_evidence_absent_from_markdown(self) -> None:
        report = compare_strategies(
            "# Notes\n\nThe reviewer checks the dates before publication.",
            [
                {"question": "What does the reviewer check?", "evidence": "checks the dates"},
                {"question": "Who approved it?", "evidence": "The board approved it"},
            ],
        )

        self.assertEqual(report["valid_case_count"], 1)
        self.assertEqual(report["invalid_case_numbers"], [2])
        self.assertEqual(len(report["strategies"]), 4)
        self.assertTrue(all(item["evidence_retention"] == 1.0 for item in report["strategies"]))

    def test_long_section_respects_maximum_chunk_size(self) -> None:
        paragraph = "A sentence about semiconductor services. " * 30
        result = chunk_markdown(f"# General\n\n{paragraph}", max_chars=300)

        self.assertGreater(len(result.chunks), 1)
        self.assertTrue(all(chunk.character_count <= 300 for chunk in result.chunks))

    def test_empty_markdown_returns_warning(self) -> None:
        result = chunk_markdown("   ")

        self.assertEqual(result.chunks, [])
        self.assertEqual(result.warnings[0].code, "empty_markdown")


if __name__ == "__main__":
    unittest.main()
