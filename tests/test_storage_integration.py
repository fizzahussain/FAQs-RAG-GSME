"""Requires a REAL isolated PostgreSQL database and a local Qdrant 1.20 service.

Set GSME_TEST_DATABASE_URL (database name must start gsme_rag_test) and
GSME_TEST_QDRANT_URL. Fake embeddings avoid external provider calls.
"""

from pathlib import Path
from tempfile import TemporaryDirectory
import os
from uuid import uuid4
from unittest import TestCase, skipUnless
from unittest.mock import patch

from gsme_rag.storage.config import Settings
from gsme_rag.storage.postgres import EvidenceStore
from gsme_rag.storage.providers import FakeEmbeddings
from gsme_rag.storage.vectors import QdrantEvidence
from gsme_rag.storage.pipeline import process_ingestion, process_index, search_evidence, reconcile

DB = os.getenv("GSME_TEST_DATABASE_URL", "")
QDRANT = os.getenv("GSME_TEST_QDRANT_URL", "")
ENABLED = bool(DB and QDRANT and "/gsme_rag_test" in DB)


@skipUnless(ENABLED, "Requires real isolated PostgreSQL and Qdrant services")
class StorageIntegrationTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(__file__).resolve().parents[1]
        self.settings = Settings(
            DB, QDRANT, Path(self.temp.name), collection="gsme_it_" + uuid4().hex[:12], dimension=32
        )
        self.store = EvidenceStore(DB)
        self.store.migrate()
        self.provider = FakeEmbeddings(32)
        self.vectors = QdrantEvidence(QDRANT, self.settings.collection, 32)
        self.tenant = "test-" + uuid4().hex
        self.other = "test-" + uuid4().hex

    def _approved_faq(self, tenant=None, document_id=None, source=None):
        tenant = tenant or self.tenant
        source = source or self.root / "examples" / "synthetic_reviewed_faq.docx"
        submitted = self.store.submit(tenant, source, self.settings.data_dir, document_id)
        if not submitted["no_op"]:
            result = process_ingestion(self.store, self.settings, submitted["job_id"])
            self.assertEqual(result["state"], "review_needed")
            inspection = self.store.inspect_version(tenant, submitted["version_id"])
            self.assertTrue(inspection["chunks"])
            self.assertTrue(
                (
                    Path(inspection["version"]["artifact_path"]).parent / "review_manifest.json"
                ).is_file()
            )
            with self.store.connect() as db:
                rows = db.execute(
                    "SELECT source_chunk_id FROM evidence_chunks WHERE version_id=%s",
                    (submitted["version_id"],),
                ).fetchall()
            self.assertEqual(len(rows), 1)
            submitted["chunk_id"] = rows[0]["source_chunk_id"]
            self.store.record_review(
                tenant,
                submitted["version_id"],
                submitted["chunk_id"],
                "approved",
                "synthetic-fixture",
                "verified generated source",
                True,
            )
            self.store.publish(
                tenant,
                submitted["document_id"],
                submitted["version_id"],
                self.provider.model_version,
                self.vectors.collection,
            )
            result = process_index(self.store, self.vectors, self.provider, self.settings,
                                   tenant, submitted["version_id"])
            self.assertEqual(result["indexed"], 1, result)
        return submitted

    def test_faq_noop_revocation_replacement_tenant_and_stale_point(self):
        first = self._approved_faq()
        with self.store.connect() as db:
            point_id = db.execute(
                "SELECT j.point_id FROM index_jobs j JOIN evidence_chunks c ON c.id=j.chunk_id "
                "WHERE c.version_id=%s", (first["version_id"],)
            ).fetchone()["point_id"]
        payload = self.vectors.existing([str(point_id)])[str(point_id)]
        self.assertEqual(set(payload), {"text", "tenant_id", "version_id",
                                        "model_version", "text_sha256"})
        question = "What is the synthetic approval deadline?"
        hits = search_evidence(
            self.store, self.vectors, self.provider, self.settings, self.tenant, question
        )
        self.assertEqual(hits[0]["chunk_id"], first["chunk_id"])
        self.assertEqual(hits[0]["source_file"], "synthetic_reviewed_faq.docx")
        self.assertEqual(
            search_evidence(
                self.store, self.vectors, self.provider, self.settings, self.other, question
            ),
            [],
        )
        calls = self.provider.calls
        rerun = self.store.submit(
            self.tenant,
            self.root / "examples" / "synthetic_reviewed_faq.docx",
            self.settings.data_dir,
            first["document_id"],
        )
        self.assertTrue(rerun["no_op"])
        self.assertEqual(
            process_index(self.store, self.vectors, self.provider, self.settings, self.tenant)[
                "indexed"
            ],
            0,
        )
        self.assertEqual(self.provider.calls, calls)
        duplicate = self._approved_faq(self.other)
        self.assertNotEqual(duplicate["document_id"], first["document_id"])
        self.store.record_review(
            self.tenant,
            first["version_id"],
            first["chunk_id"],
            "rejected",
            "synthetic-fixture",
            "revoked for test",
            True,
        )
        self.assertEqual(
            search_evidence(
                self.store, self.vectors, self.provider, self.settings, self.tenant, question
            ),
            [],
        )
        self.assertGreaterEqual(reconcile(self.store, self.vectors, self.provider.model_version), 1)
        # Version replacement: old Qdrant point still exists until reconciliation.
        from docx import Document

        changed = Path(self.temp.name) / "replacement.docx"
        d = Document()
        d.add_heading("General", 1)
        q = d.add_paragraph()
        q.add_run("What is the replacement value?").bold = True
        d.add_paragraph("The replacement value is synthetic revision two.")
        d.save(changed)
        newer = self._approved_faq(source=changed, document_id=first["document_id"])
        self.assertNotEqual(newer["version_id"], first["version_id"])
        self.assertEqual(
            search_evidence(
                self.store,
                self.vectors,
                self.provider,
                self.settings,
                self.tenant,
                "What is the replacement value?",
            )[0]["version_id"],
            newer["version_id"],
        )
        self.assertEqual(
            search_evidence(
                self.store, self.vectors, self.provider, self.settings, self.tenant, question
            )[0]["version_id"],
            newer["version_id"],
        )
        self.assertGreaterEqual(reconcile(self.store, self.vectors, self.provider.model_version), 0)

    def test_full_shared_schedule_typed_exact_and_tenant(self):
        header = "Tech (nm),Flavor,Trial,Final,Tapeout,Est. Ship\n"
        row = "28,HPC+ RF,10/01/2026,10/15/2026,11/01/2026,12/01/2026\n"
        for kind in ("full", "shared"):
            source = Path(self.temp.name) / (kind + "_schedules.csv")
            source.write_text(header + row, encoding="utf-8")
            submitted = self.store.submit(self.tenant, source, self.settings.data_dir)
            processed = process_ingestion(self.store, self.settings, submitted["job_id"])
            self.assertEqual(processed["schedule_rows"], 1)
            found = self.store.schedules(
                self.tenant,
                kind,
                technology_nm=28,
                flavor="HPC+ RF",
                document_id=submitted["document_id"],
            )
            self.assertEqual(len(found), 1)
            self.assertEqual(found[0]["source_row_number"], 2)
            self.assertEqual(self.store.schedules(self.other, kind), [])
        self.assertEqual(self.store.schedules(self.tenant, "full", technology_nm=14), [])

    def test_invalid_source_reports_operation_without_content(self):
        source = Path(self.temp.name) / "invalid.docx"
        source.write_bytes(b"private test bytes, not a DOCX")
        submitted = self.store.submit(self.tenant, source, self.settings.data_dir)
        result = process_ingestion(self.store, self.settings, submitted["job_id"])
        self.assertEqual(result["state"], "failed")
        status = self.store.job_status(self.tenant, submitted["job_id"])
        self.assertIn("ingestion parse failed", status["last_error"])
        self.assertNotIn("private test bytes", status["last_error"])

    def test_pending_and_rejected_chunks_never_reach_embedding_provider(self):
        source = self.root / "examples" / "synthetic_reviewed_faq.docx"
        submitted = self.store.submit(self.tenant, source, self.settings.data_dir)
        self.assertEqual(process_ingestion(self.store, self.settings, submitted["job_id"])["state"],
                         "review_needed")
        with self.assertRaisesRegex(ValueError, "source-verified approved"):
            self.store.publish(self.tenant, submitted["document_id"],
                               submitted["version_id"], self.provider.model_version,
                               self.vectors.collection)
        self.assertEqual(process_index(self.store, self.vectors, self.provider,
                                       self.settings, self.tenant)["indexed"], 0)
        self.assertEqual(self.provider.calls, 0)

        with self.store.connect() as db:
            chunk_id = db.execute("SELECT source_chunk_id FROM evidence_chunks WHERE version_id=%s",
                                  (submitted["version_id"],)).fetchone()["source_chunk_id"]
        self.store.record_review(self.tenant, submitted["version_id"], chunk_id,
                                 "rejected", "synthetic-fixture", "rejected for test", True)
        self.assertEqual(process_index(self.store, self.vectors, self.provider,
                                       self.settings, self.tenant)["indexed"], 0)
        self.assertEqual(self.provider.calls, 0)

    def test_source_order_pagination_and_immutable_reviewed_evidence(self):
        from docx import Document
        import psycopg

        source = Path(self.temp.name) / "many.docx"
        document = Document()
        for index in range(105):
            document.add_heading(f"What is source item {index}?", level=2)
            document.add_paragraph(f"The source answer is item {index}.")
        document.save(source)
        submitted = self.store.submit(self.tenant, source, self.settings.data_dir)
        self.assertEqual(process_ingestion(self.store, self.settings, submitted["job_id"])["chunks"], 105)
        first = self.store.inspect_version(self.tenant, submitted["version_id"], 100)
        last = self.store.inspect_version(self.tenant, submitted["version_id"], 100, first["next_offset"])
        self.assertEqual([chunk["chunk_index"] for chunk in first["chunks"] + last["chunks"]], list(range(105)))
        self.assertEqual(first["total_chunks"], 105)
        self.assertIsNone(last["next_offset"])
        with self.assertRaises(psycopg.errors.RaiseException):
            with self.store.connect() as db:
                db.execute("UPDATE evidence_chunks SET evidence_text='changed' WHERE version_id=%s",
                           (submitted["version_id"],))

    def test_bad_candidate_does_not_block_healthy_indexing(self):
        from dataclasses import replace
        from docx import Document

        source = Path(self.temp.name) / "mixed.docx"
        document = Document()
        document.add_heading("What is the long answer?", 2)
        document.add_paragraph("synthetic lengthy detail " * 30)
        document.add_heading("What is short?", 2)
        document.add_paragraph("A short answer.")
        document.save(source)
        submitted = self.store.submit(self.tenant, source, self.settings.data_dir)
        process_ingestion(self.store, self.settings, submitted["job_id"])
        chunks = self.store.inspect_version(self.tenant, submitted["version_id"])["chunks"]
        for chunk in chunks:
            self.store.record_review(self.tenant, submitted["version_id"], chunk["source_chunk_id"],
                                     "approved", "synthetic", "Compared generated source", True)
        self.store.publish(self.tenant, submitted["document_id"], submitted["version_id"],
                           self.provider.model_version, self.vectors.collection)
        result = process_index(self.store, self.vectors, self.provider,
                               replace(self.settings, max_tokens=30), self.tenant, submitted["version_id"])
        self.assertEqual((result["indexed"], result["failed"]), (1, 1), result)
        with self.store.connect() as db:
            bad = db.execute("SELECT retryable,last_error FROM index_jobs WHERE state='failed' AND chunk_id IN "
                             "(SELECT id FROM evidence_chunks WHERE version_id=%s)",
                             (submitted["version_id"],)).fetchone()
        self.assertFalse(bad["retryable"])
        self.assertIn("budget", bad["last_error"])
        self.assertEqual(process_index(self.store, self.vectors, self.provider,
                                      self.settings, self.tenant, submitted["version_id"])["indexed"], 0)

    def test_generic_faq_csv_is_a_document_and_parser_upgrade_creates_new_version(self):
        source = Path(self.temp.name) / "questions.csv"
        source.write_text("Question,Answer\nWhat is Muse?,Semiconductor services.\n", encoding="utf-8")
        first = self.store.submit(self.tenant, source, self.settings.data_dir)
        self.assertEqual(process_ingestion(self.store, self.settings, first["job_id"])["state"], "review_needed")
        with patch("gsme_rag.storage.postgres.PARSER_VERSION", "next-parser-version"):
            second = self.store.submit(self.tenant, source, self.settings.data_dir, first["document_id"])
        self.assertFalse(second["no_op"])
        self.assertNotEqual(first["version_id"], second["version_id"])

    def test_missing_derived_point_can_be_repaired_without_changing_approval(self):
        from gsme_rag.storage.pipeline import repair_missing_vectors
        first = self._approved_faq()
        point = self.store.approved_point(self.tenant, first["version_id"], first["chunk_id"],
                                         self.provider.model_version, self.vectors.collection)
        self.vectors.delete([str(point["point_id"])])
        self.assertEqual(repair_missing_vectors(self.store, self.vectors, self.provider.model_version,
                                                self.tenant, first["version_id"]), 1)
        result = process_index(self.store, self.vectors, self.provider, self.settings,
                               self.tenant, first["version_id"])
        self.assertEqual(result["indexed"], 1)
        self.assertIn(str(point["point_id"]), self.vectors.existing([str(point["point_id"])]))

    def test_schedule_write_failure_does_not_partially_publish(self):
        source = Path(self.temp.name) / "full_schedules.csv"
        source.write_text(
            "Tech (nm),Flavor,Trial,Final,Tapeout,Est. Ship\n"
            "28,RF,10/01/2026,10/15/2026,11/01/2026,12/01/2026\n",
            encoding="utf-8",
        )
        submitted = self.store.submit(self.tenant, source, self.settings.data_dir)
        with patch.object(EvidenceStore, "_publish", side_effect=RuntimeError("private detail")):
            result = process_ingestion(self.store, self.settings, submitted["job_id"])
        self.assertEqual(result["state"], "failed")
        self.assertEqual(result["error"], "RuntimeError: ingestion database_write failed")
        with self.store.connect() as db:
            count = db.execute("SELECT count(*) FROM schedule_rows WHERE version_id=%s",
                               (submitted["version_id"],)).fetchone()["count"]
            current = db.execute("SELECT current_version_id FROM documents WHERE id=%s",
                                 (submitted["document_id"],)).fetchone()["current_version_id"]
        self.assertEqual(count, 0)
        self.assertIsNone(current)

    def test_expired_worker_attempt_cannot_complete_newer_claim(self):
        source = self.root / "examples" / "synthetic_reviewed_faq.docx"
        submitted = self.store.submit(self.tenant, source, self.settings.data_dir)
        old = self.store.claim_ingestion(30, submitted["job_id"])
        with self.store.connect() as db:
            db.execute("UPDATE ingestion_jobs SET claimed_at=now()-interval '1 hour' WHERE id=%s",
                       (submitted["job_id"],))
        new = self.store.claim_ingestion(30, submitted["job_id"])
        self.assertEqual(new["attempts"], old["attempts"] + 1)
        self.assertFalse(self.store.heartbeat_ingestion(old["id"], old["attempts"]))
        with self.assertRaisesRegex(ValueError, "lease changed"):
            self.store.complete_ingestion(old, None, [], [], "")
        self.assertEqual(self.store.job_status(self.tenant, submitted["job_id"])["state"],
                         "running")

    def test_superseded_qdrant_point_is_filtered_before_cleanup(self):
        first = self._approved_faq()
        with self.store.connect() as db:
            old_point = db.execute(
                """SELECT j.point_id FROM index_jobs j JOIN evidence_chunks c
                ON c.id=j.chunk_id WHERE c.version_id=%s""",
                (first["version_id"],),
            ).fetchone()["point_id"]
        from docx import Document

        changed = Path(self.temp.name) / "replacement.docx"
        d = Document()
        d.add_heading("General", 1)
        q = d.add_paragraph()
        q.add_run("What is the replacement value?").bold = True
        d.add_paragraph("The replacement value is synthetic revision two.")
        d.save(changed)
        newer = self._approved_faq(source=changed, document_id=first["document_id"])
        self.assertIn(str(old_point), self.vectors.existing([str(old_point)]))
        hits = search_evidence(
            self.store,
            self.vectors,
            self.provider,
            self.settings,
            self.tenant,
            "What is the synthetic approval deadline?",
        )
        self.assertTrue(all(h["version_id"] == newer["version_id"] for h in hits))
        self.assertGreaterEqual(reconcile(self.store, self.vectors, self.provider.model_version), 1)
